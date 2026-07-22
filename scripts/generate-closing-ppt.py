#!/usr/bin/env python3
"""Generate visually enhanced closing presentation PPT for CAT Apps."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Brand palette
PRIMARY = RGBColor(0x12, 0x3B, 0x8F)
PRIMARY_LIGHT = RGBColor(0x1E, 0x56, 0xC4)
PRIMARY_DARK = RGBColor(0x0A, 0x24, 0x5C)
ACCENT = RGBColor(0xF5, 0xA6, 0x23)
ACCENT_LIGHT = RGBColor(0xFF, 0xD9, 0x6A)
SUCCESS = RGBColor(0x16, 0xA3, 0x4A)
SUCCESS_BG = RGBColor(0xDC, 0xFC, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
LIGHT_BG = RGBColor(0xF0, 0xF4, 0xFA)
SOFT_BLUE = RGBColor(0xE8, 0xF0, 0xFE)
SOFT_GOLD = RGBColor(0xFF, 0xF7, 0xED)

ROLE_COLORS = {
    "admin": RGBColor(0x7C, 0x3A, 0xED),
    "mentor": RGBColor(0x0E, 0xA5, 0xE9),
    "peserta": RGBColor(0x16, 0xA3, 0x4A),
    "parent": RGBColor(0xF5, 0xA6, 0x23),
}

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Presentasi-Closing-CAT-Apps-v2.pptx"
LOGO_PATH = ROOT / "resources" / "assets" / "logo.png"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, shape_type, left, top, width, height, fill, line=None, radius=None):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    if radius is not None and hasattr(shape, "adjustments") and len(shape.adjustments) > 0:
        shape.adjustments[0] = radius
    return shape


def set_text(shape, text, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return tf


def add_footer(slide, prs, page: int, total: int) -> None:
    bar = add_shape(slide, MSO_SHAPE.RECTANGLE, 0, SLIDE_H - Inches(0.38), SLIDE_W, Inches(0.38), PRIMARY_DARK)
    left = slide.shapes.add_textbox(Inches(0.5), SLIDE_H - Inches(0.34), Inches(5), Inches(0.28))
    lp = left.text_frame.paragraphs[0]
    lp.text = "Pratistha Cendekia Prestasi · CAT Apps"
    lp.font.size = Pt(9)
    lp.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)

    right = slide.shapes.add_textbox(Inches(11.5), SLIDE_H - Inches(0.34), Inches(1.5), Inches(0.28))
    rp = right.text_frame.paragraphs[0]
    rp.text = f"{page} / {total}"
    rp.font.size = Pt(9)
    rp.font.color.rgb = ACCENT_LIGHT
    rp.alignment = PP_ALIGN.RIGHT


def add_content_header(slide, title: str, subtitle: str | None = None, badge: str | None = None) -> None:
    set_bg(slide, LIGHT_BG)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.35), PRIMARY)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(1.35), SLIDE_W, Inches(0.06), ACCENT)

    if badge:
        badge_shape = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(0.18), Inches(1.1), Inches(0.32), ACCENT, radius=0.25
        )
        set_text(badge_shape, badge, size=11, bold=True, color=PRIMARY_DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.52 if badge else 0.35), Inches(11.5), Inches(0.55))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(30)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.55), Inches(0.98), Inches(11.5), Inches(0.3))
        sp = sub.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(13)
        sp.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)


def add_section_divider(prs, blank, number: str, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PRIMARY_DARK)
    add_shape(slide, MSO_SHAPE.OVAL, Inches(-1.2), Inches(-1.2), Inches(4.5), Inches(4.5), PRIMARY_LIGHT)
    add_shape(slide, MSO_SHAPE.OVAL, Inches(10.5), Inches(5.2), Inches(3.5), Inches(3.5), RGBColor(0x0F, 0x33, 0x7A))
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(6.85), SLIDE_W, Inches(0.65), ACCENT)

    num = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(3), Inches(1.2))
    np = num.text_frame.paragraphs[0]
    np.text = number
    np.font.size = Pt(72)
    np.font.bold = True
    np.font.color.rgb = ACCENT

    ttl = slide.shapes.add_textbox(Inches(0.9), Inches(3.2), Inches(11), Inches(0.9))
    tp = ttl.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(36)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    sub = slide.shapes.add_textbox(Inches(0.9), Inches(4.15), Inches(10.5), Inches(0.6))
    sp = sub.text_frame.paragraphs[0]
    sp.text = subtitle
    sp.font.size = Pt(18)
    sp.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)


def add_stat_cards(slide, cards: list[tuple[str, str, str]]) -> None:
    """cards: (big_number, label, sublabel)"""
    n = len(cards)
    gap = 0.25
    width = (12.0 - gap * (n - 1)) / n
    for i, (num, label, sub) in enumerate(cards):
        x = 0.65 + i * (width + gap)
        card = add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(1.65),
            Inches(width),
            Inches(1.55),
            WHITE,
            PRIMARY_LIGHT,
            radius=0.12,
        )
        accent = add_shape(
            slide, MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.65), Inches(width), Inches(0.08), ACCENT
        )
        accent.line.fill.background()

        num_box = slide.shapes.add_textbox(Inches(x + 0.15), Inches(1.85), Inches(width - 0.3), Inches(0.55))
        np = num_box.text_frame.paragraphs[0]
        np.text = num
        np.font.size = Pt(34)
        np.font.bold = True
        np.font.color.rgb = PRIMARY
        np.alignment = PP_ALIGN.CENTER

        lbl = slide.shapes.add_textbox(Inches(x + 0.1), Inches(2.45), Inches(width - 0.2), Inches(0.35))
        lp = lbl.text_frame.paragraphs[0]
        lp.text = label
        lp.font.size = Pt(13)
        lp.font.bold = True
        lp.font.color.rgb = TEXT
        lp.alignment = PP_ALIGN.CENTER

        if sub:
            sb = slide.shapes.add_textbox(Inches(x + 0.1), Inches(2.78), Inches(width - 0.2), Inches(0.3))
            sp = sb.text_frame.paragraphs[0]
            sp.text = sub
            sp.font.size = Pt(10)
            sp.font.color.rgb = MUTED
            sp.alignment = PP_ALIGN.CENTER


def add_icon_bullets(slide, items: list[tuple[str, str]], left=0.65, top=1.55, col_width=5.9, cols=2):
    """items: (icon_emoji, text)"""
    per_col = (len(items) + cols - 1) // cols
    for idx, (icon, text) in enumerate(items):
        col = idx // per_col
        row = idx % per_col
        x = left + col * (col_width + 0.35)
        y = top + row * 0.95

        card = add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(col_width),
            Inches(0.78),
            WHITE,
            RGBColor(0xE5, 0xE7, 0xEB),
            radius=0.15,
        )
        dot = add_shape(slide, MSO_SHAPE.OVAL, Inches(x + 0.12), Inches(y + 0.18), Inches(0.42), Inches(0.42), SOFT_BLUE)
        set_text(dot, icon, size=16, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        tb = slide.shapes.add_textbox(Inches(x + 0.65), Inches(y + 0.12), Inches(col_width - 0.75), Inches(0.55))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT


def add_flow_pipeline(slide, steps: list[str]) -> None:
    box_w = 11.5
    box_h = 0.62
    start_x = 0.9
    start_y = 1.55
    gap = 0.22

    colors = [PRIMARY, PRIMARY_LIGHT, RGBColor(0x25, 0x63, 0xEB), RGBColor(0x0E, 0xA5, 0xE9), RGBColor(0x16, 0xA3, 0x4A), ACCENT]

    for i, step in enumerate(steps):
        y = start_y + i * (box_h + gap)
        color = colors[i % len(colors)]

        box = add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(start_x),
            Inches(y),
            Inches(box_w),
            Inches(box_h),
            color,
            radius=0.1,
        )

        num = add_shape(
            slide, MSO_SHAPE.OVAL, Inches(start_x + 0.12), Inches(y + 0.1), Inches(0.42), Inches(0.42), WHITE
        )
        set_text(num, str(i + 1), size=14, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        tb = slide.shapes.add_textbox(Inches(start_x + 0.65), Inches(y + 0.1), Inches(box_w - 0.8), Inches(0.45))
        p = tb.text_frame.paragraphs[0]
        p.text = step
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE

        if i < len(steps) - 1:
            arrow = add_shape(
                slide,
                MSO_SHAPE.DOWN_ARROW,
                Inches(start_x + box_w / 2 - 0.15),
                Inches(y + box_h + 0.02),
                Inches(0.3),
                Inches(gap - 0.04),
                ACCENT,
            )


def add_role_cards(slide) -> None:
    roles = [
        ("Admin", "Kelola platform, user, sertifikat, review pendaftaran", ROLE_COLORS["admin"], "A"),
        ("Mentor", "Kelas, bank soal, nilai jasmani, laporan harian", ROLE_COLORS["mentor"], "M"),
        ("Peserta", "Pendaftaran, kelas, quiz, ujian, dashboard", ROLE_COLORS["peserta"], "P"),
        ("Orang Tua", "Pantau anak, laporan mingguan, unduh PDF", ROLE_COLORS["parent"], "O"),
    ]
    positions = [(0.65, 1.55), (6.75, 1.55), (0.65, 4.05), (6.75, 4.05)]

    for (name, desc, color, initial), (x, y) in zip(roles, positions):
        card = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.9), Inches(2.2), WHITE, color, radius=0.08
        )
        header = add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(5.9), Inches(0.55), color)
        header.line.fill.background()

        avatar = add_shape(
            slide, MSO_SHAPE.OVAL, Inches(x + 0.2), Inches(y + 0.08), Inches(0.4), Inches(0.4), WHITE
        )
        set_text(avatar, initial, size=14, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        title = slide.shapes.add_textbox(Inches(x + 0.75), Inches(y + 0.12), Inches(4.8), Inches(0.35))
        tp = title.text_frame.paragraphs[0]
        tp.text = name
        tp.font.size = Pt(18)
        tp.font.bold = True
        tp.font.color.rgb = WHITE

        body = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.75), Inches(5.4), Inches(1.2))
        bp = body.text_frame.paragraphs[0]
        bp.text = desc
        bp.font.size = Pt(14)
        bp.font.color.rgb = TEXT


def add_tech_cards(slide) -> None:
    stack = [
        ("Backend", "Laravel 12\nPHP 8.2+", PRIMARY),
        ("Frontend", "Vue 3 · Pinia\nTailwind 4 · PWA", PRIMARY_LIGHT),
        ("Database", "PostgreSQL\nSQLite (dev)", RGBColor(0x0E, 0xA5, 0xE9)),
        ("Storage", "Cloudflare R2\nBerkas privat", RGBColor(0x16, 0xA3, 0x4A)),
        ("Deploy", "Docker · Coolify\nApache · Backup", RGBColor(0x7C, 0x3A, 0xED)),
        ("Keamanan", "CSRF · Role auth\nEmail verify", ACCENT),
    ]
    cols, rows = 3, 2
    w, h = 3.85, 1.85
    gap_x, gap_y = 0.3, 0.28
    start_x, start_y = 0.65, 1.55

    for i, (title, desc, color) in enumerate(stack):
        col, row = i % cols, i // cols
        x = start_x + col * (w + gap_x)
        y = start_y + row * (h + gap_y)

        card = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h), WHITE, color, radius=0.1
        )
        stripe = add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.12), Inches(h), color)
        stripe.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(x + 0.3), Inches(y + 0.25), Inches(w - 0.45), Inches(h - 0.35))
        tf = tb.text_frame
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT
        p2.space_before = Pt(6)


def add_timeline_visual(slide) -> None:
    milestones = [
        ("4 Mei", "Kick-off", "Mulai pengembangan"),
        ("Minggu 3", "90%", "Modul inti selesai"),
        ("31 Mei", "Go Live", "Deploy produksi"),
        ("8 Juli", "Closing", "Serah terima"),
        ("Mei '27", "Support", "Maintenance 12 bln"),
    ]
    n = len(milestones)
    line_y = 3.2
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(line_y), Inches(11.5), Inches(0.06), ACCENT)

    step_w = 11.5 / n
    for i, (date, title, sub) in enumerate(milestones):
        cx = 0.9 + step_w * i + step_w / 2
        dot = add_shape(
            slide, MSO_SHAPE.OVAL, Inches(cx - 0.18), Inches(line_y - 0.15), Inches(0.36), Inches(0.36), PRIMARY
        )
        dot.line.color.rgb = ACCENT
        dot.line.width = Pt(2)

        date_box = slide.shapes.add_textbox(Inches(cx - 0.7), Inches(2.35), Inches(1.4), Inches(0.35))
        dp = date_box.text_frame.paragraphs[0]
        dp.text = date
        dp.font.size = Pt(13)
        dp.font.bold = True
        dp.font.color.rgb = ACCENT
        dp.alignment = PP_ALIGN.CENTER

        title_box = slide.shapes.add_textbox(Inches(cx - 0.95), Inches(3.65), Inches(1.9), Inches(0.35))
        tp = title_box.text_frame.paragraphs[0]
        tp.text = title
        tp.font.size = Pt(15)
        tp.font.bold = True
        tp.font.color.rgb = PRIMARY
        tp.alignment = PP_ALIGN.CENTER

        sub_box = slide.shapes.add_textbox(Inches(cx - 1.0), Inches(4.05), Inches(2.0), Inches(0.55))
        sp = sub_box.text_frame.paragraphs[0]
        sp.text = sub
        sp.font.size = Pt(10)
        sp.font.color.rgb = MUTED
        sp.alignment = PP_ALIGN.CENTER


def add_module_grid(slide, modules: list[tuple[str, str, str]]) -> None:
    """(number, title, status)"""
    cols = 4
    w, h = 2.85, 1.15
    gap = 0.22
    start_x, start_y = 0.65, 1.55

    for i, (num, title, status) in enumerate(modules):
        col, row = i % cols, i // cols
        x = start_x + col * (w + gap)
        y = start_y + row * (h + gap)

        card = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h), WHITE, PRIMARY_LIGHT, radius=0.1
        )

        num_shape = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.12), Inches(y + 0.12), Inches(0.45), Inches(0.45), PRIMARY, radius=0.2
        )
        set_text(num_shape, num, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        title_box = slide.shapes.add_textbox(Inches(x + 0.65), Inches(y + 0.15), Inches(w - 0.75), Inches(0.45))
        tp = title_box.text_frame.paragraphs[0]
        tp.text = title
        tp.font.size = Pt(11)
        tp.font.bold = True
        tp.font.color.rgb = TEXT

        badge_color = SUCCESS_BG if "Selesai" in status else SOFT_GOLD
        badge_text = SUCCESS if "Selesai" in status else RGBColor(0xB4, 0x53, 0x09)
        badge = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.12), Inches(y + 0.72), Inches(1.0), Inches(0.28), badge_color, radius=0.3
        )
        set_text(badge, status.replace("*", ""), size=9, bold=True, color=badge_text, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def add_feature_pills(slide, features: list[tuple[str, str]]) -> None:
    cols = 2
    w, h = 5.9, 0.95
    gap_x, gap_y = 0.35, 0.22
    start_x, start_y = 0.65, 1.55

    for i, (title, desc) in enumerate(features):
        col, row = i % cols, i // cols
        x = start_x + col * (w + gap_x)
        y = start_y + row * (h + gap_y)

        card = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h), WHITE, ACCENT, radius=0.12
        )
        pill = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.15), Inches(y + 0.15), Inches(1.55), Inches(0.3), SOFT_GOLD, radius=0.4
        )
        set_text(pill, title, size=10, bold=True, color=RGBColor(0xB4, 0x53, 0x09), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        desc_box = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.5), Inches(w - 0.3), Inches(0.38))
        dp = desc_box.text_frame.paragraphs[0]
        dp.text = desc
        dp.font.size = Pt(12)
        dp.font.color.rgb = TEXT


def add_identity_cards(slide) -> None:
    parties = [
        ("PIHAK PERTAMA", "Developer", "Lulu Maulana Malik", "Sukasari No. 34, Cibiru, Bandung", PRIMARY),
        ("PIHAK KEDUA", "Klien", "BJP. (P) Drs. H. Awang Anwarudin, M.H.", "Kantor PP Polri Daerah Jabar, Bandung", PRIMARY_LIGHT),
    ]
    for i, (label, role, name, addr, color) in enumerate(parties):
        x = 0.65 + i * 6.1
        card = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.65), Inches(5.85), Inches(4.5), WHITE, color, radius=0.08
        )
        header = add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.65), Inches(5.85), Inches(0.9), color)
        header.line.fill.background()

        lbl = slide.shapes.add_textbox(Inches(x + 0.3), Inches(1.78), Inches(5.2), Inches(0.3))
        lp = lbl.text_frame.paragraphs[0]
        lp.text = label
        lp.font.size = Pt(11)
        lp.font.bold = True
        lp.font.color.rgb = ACCENT_LIGHT

        rl = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2.08), Inches(5.2), Inches(0.35))
        rp = rl.text_frame.paragraphs[0]
        rp.text = role
        rp.font.size = Pt(20)
        rp.font.bold = True
        rp.font.color.rgb = WHITE

        nm = slide.shapes.add_textbox(Inches(x + 0.35), Inches(2.85), Inches(5.1), Inches(0.8))
        np = nm.text_frame.paragraphs[0]
        np.text = name
        np.font.size = Pt(16)
        np.font.bold = True
        np.font.color.rgb = TEXT

        ad = slide.shapes.add_textbox(Inches(x + 0.35), Inches(3.75), Inches(5.1), Inches(1.2))
        ap = ad.text_frame.paragraphs[0]
        ap.text = addr
        ap.font.size = Pt(13)
        ap.font.color.rgb = MUTED


def add_checklist(slide, items: list[str], top=3.45) -> None:
    for i, item in enumerate(items):
        y = top + i * 0.52
        check = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(y), Inches(0.28), Inches(0.28), SUCCESS_BG, SUCCESS, radius=0.2
        )
        set_text(check, "✓", size=12, bold=True, color=SUCCESS, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        tb = slide.shapes.add_textbox(Inches(1.15), Inches(y - 0.02), Inches(11.5), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = item
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT


def add_cover(prs, blank) -> None:
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PRIMARY_DARK)

    add_shape(slide, MSO_SHAPE.OVAL, Inches(-2), Inches(4.5), Inches(6), Inches(6), PRIMARY_LIGHT)
    add_shape(slide, MSO_SHAPE.OVAL, Inches(9.5), Inches(-2.5), Inches(5.5), Inches(5.5), RGBColor(0x0F, 0x33, 0x7A))
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(6.75), SLIDE_W, Inches(0.75), ACCENT)

    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(5.65), Inches(0.55), height=Inches(1.35))

    badge = add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.85), Inches(2.05), Inches(3.6), Inches(0.38), ACCENT, radius=0.4
    )
    set_text(badge, "PRESENTASI CLOSING · JULI 2026", size=11, bold=True, color=PRIMARY_DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    title = slide.shapes.add_textbox(Inches(0.8), Inches(2.65), Inches(11.7), Inches(1.1))
    tp = title.text_frame.paragraphs[0]
    tp.text = "Pratistha Cendekia Prestasi"
    tp.font.size = Pt(46)
    tp.font.bold = True
    tp.font.color.rgb = WHITE
    tp.alignment = PP_ALIGN.CENTER

    sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.65), Inches(11.7), Inches(0.55))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "CAT Apps — Platform Bimbel & Simulasi Seleksi AKPOL"
    sp.font.size = Pt(20)
    sp.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)
    sp.alignment = PP_ALIGN.CENTER

    meta_card = add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.2), Inches(4.55), Inches(6.9), Inches(1.55), RGBColor(0x0F, 0x33, 0x7A), ACCENT, radius=0.08
    )
    meta = slide.shapes.add_textbox(Inches(3.45), Inches(4.75), Inches(6.4), Inches(1.2))
    mf = meta.text_frame
    mf.text = "pratisthaindonesia.com\nPeriode: 4 Mei – 31 Mei 2026 · Bandung"
    for p in mf.paragraphs:
        p.font.size = Pt(14)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER


def add_closing(prs, blank) -> None:
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PRIMARY_DARK)
    add_shape(slide, MSO_SHAPE.OVAL, Inches(8.5), Inches(-1.5), Inches(5), Inches(5), PRIMARY_LIGHT)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(6.75), SLIDE_W, Inches(0.75), ACCENT)

    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(5.65), Inches(0.45), height=Inches(1.1))

    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.75), Inches(11.7), Inches(0.8))
    tp = title.text_frame.paragraphs[0]
    tp.text = "Penutup & Serah Terima"
    tp.font.size = Pt(38)
    tp.font.bold = True
    tp.font.color.rgb = WHITE
    tp.alignment = PP_ALIGN.CENTER

    card = add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(2.85), Inches(8.9), Inches(2.0), RGBColor(0x0F, 0x33, 0x7A), ACCENT, radius=0.06
    )
    body = slide.shapes.add_textbox(Inches(2.5), Inches(3.1), Inches(8.3), Inches(1.5))
    bf = body.text_frame
    bf.text = (
        "Aplikasi CAT Apps telah selesai dikembangkan\n"
        "dan beroperasi di produksi.\n\n"
        "Siap mendukung operasional bimbel AKPOL."
    )
    for p in bf.paragraphs:
        p.font.size = Pt(18)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    contact = slide.shapes.add_textbox(Inches(1.5), Inches(5.15), Inches(10.3), Inches(0.9))
    cf = contact.text_frame
    cf.text = "pratisthaindonesia.com  ·  WA +62 813-8964-488  ·  @pratistha.cendikia"
    cp = cf.paragraphs[0]
    cp.font.size = Pt(13)
    cp.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)
    cp.alignment = PP_ALIGN.CENTER

    thanks = slide.shapes.add_textbox(Inches(0.8), Inches(5.95), Inches(11.7), Inches(0.55))
    tp2 = thanks.text_frame.paragraphs[0]
    tp2.text = "Terima Kasih"
    tp2.font.size = Pt(32)
    tp2.font.bold = True
    tp2.font.color.rgb = ACCENT
    tp2.alignment = PP_ALIGN.CENTER


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    total_slides = 22  # approximate for footer

    # Cover
    add_cover(prs, blank)

    # Agenda
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Agenda Presentasi", "Roadmap sesi dari latar belakang hingga penutup")
    add_icon_bullets(
        slide,
        [
            ("01", "Latar belakang & identitas proyek"),
            ("02", "Ringkasan eksekutif & statistik"),
            ("03", "Alur platform end-to-end"),
            ("04", "Peran pengguna & tech stack"),
            ("05", "12 modul ruang lingkup pekerjaan"),
            ("06", "Fitur tambahan & timeline"),
            ("07", "Demo alur & portal orang tua"),
            ("08", "Catatan teknis & penutup"),
        ],
        top=1.55,
        col_width=5.9,
    )
    add_footer(slide, prs, 2, total_slides)

    add_section_divider(prs, blank, "01", "Perkenalan Proyek", "Identitas, latar belakang, dan ringkasan hasil")

    # Identitas
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Identitas Proyek", "Perjanjian Kerja Sama — 4 Mei 2026")
    add_identity_cards(slide)
    add_footer(slide, prs, 4, total_slides)

    # Latar belakang
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Latar Belakang & Tujuan")
    left_card = add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.55), Inches(5.9), Inches(5.0), WHITE, PRIMARY, radius=0.08
    )
    set_text(left_card, "", size=1)
    lh = slide.shapes.add_textbox(Inches(0.95), Inches(1.85), Inches(5.3), Inches(4.3))
    ltf = lh.text_frame
    ltf.text = "Mengapa CAT Apps?"
    ltf.paragraphs[0].font.size = Pt(20)
    ltf.paragraphs[0].font.bold = True
    ltf.paragraphs[0].font.color.rgb = PRIMARY
    for line in [
        "Bimbel persiapan seleksi AKPOL membutuhkan sistem terintegrasi.",
        "Pendaftaran, kelas, ujian CAT, dan monitoring harus dalam satu platform.",
        "CAT Apps mendigitalisasi seluruh alur operasional bimbel.",
    ]:
        p = ltf.add_paragraph()
        p.text = line
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT
        p.space_before = Pt(10)

    goals = [
        ("Centralized", "Manajemen soal, ujian & kelas"),
        ("Real-time", "Dashboard & laporan perkembangan"),
        ("Connected", "Peserta · Mentor · Admin · Ortu"),
        ("Public", "Tryout gratis sebagai funnel"),
    ]
    for i, (gtitle, gdesc) in enumerate(goals):
        y = 1.55 + i * 1.22
        card = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(y), Inches(5.9), Inches(1.05), SOFT_BLUE, PRIMARY_LIGHT, radius=0.1
        )
        tb = slide.shapes.add_textbox(Inches(7.0), Inches(y + 0.15), Inches(5.4), Inches(0.75))
        tf = tb.text_frame
        p1 = tf.paragraphs[0]
        p1.text = gtitle
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY
        p2 = tf.add_paragraph()
        p2.text = gdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT
    add_footer(slide, prs, 5, total_slides)

    # Ringkasan
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Ringkasan Eksekutif", "Status penyelesaian proyek")
    add_stat_cards(
        slide,
        [
            ("12/12", "Modul Selesai", "Ruang lingkup PK"),
            ("4", "Peran User", "Admin · Mentor · Peserta · Ortu"),
            ("Live", "Produksi", "pratisthaindonesia.com"),
            ("12 bln", "Maintenance", "s.d. Mei 2027"),
        ],
    )
    add_checklist(
        slide,
        [
            "Arsitektur modern Laravel 12 + Vue 3 + PostgreSQL + Docker",
            "Fitur bonus: portal orang tua, sertifikat, PWA, anti-cheat, multi-bahasa",
            "Semua modul inti telah di-deploy dan siap operasional",
        ],
    )
    add_footer(slide, prs, 6, total_slides)

    add_section_divider(prs, blank, "02", "Alur & Peran", "Bagaimana platform bekerja dari awal sampai akhir")

    # Flow
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Alur Platform", "Perjalanan peserta dari beranda hingga laporan orang tua")
    add_flow_pipeline(
        slide,
        [
            "Beranda Publik + Tryout Gratis",
            "Daftar Akun → Verifikasi Email → Pendaftaran Multi-Tahap",
            "Approval Admin → Dashboard Terbuka",
            "Kelas Saya (materi & quiz) + Ujian Lintas Mapel",
            "Staff Input Nilai Jasmani & Laporan Harian",
            "Orang Tua Pantau Perkembangan & Unduh PDF",
        ],
    )
    add_footer(slide, prs, 8, total_slides)

    # Roles
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Peran Pengguna", "Empat role dengan hak akses berbeda")
    add_role_cards(slide)
    add_footer(slide, prs, 9, total_slides)

    # Tech
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Spesifikasi Teknis", "Arsitektur modern & scalable")
    add_tech_cards(slide)
    add_footer(slide, prs, 10, total_slides)

    add_section_divider(prs, blank, "03", "Modul Aplikasi", "12 poin ruang lingkup pekerjaan — status selesai")

    # 12 modules grid
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Ruang Lingkup — 12 Modul", "Semua modul inti telah diselesaikan")
    add_module_grid(
        slide,
        [
            ("1", "Manajemen Soal", "Selesai"),
            ("2", "Ujian Online", "Selesai"),
            ("3", "Pendaftaran", "Selesai"),
            ("4", "Tryout Gratis", "Selesai"),
            ("5", "Penilaian Otomatis", "Selesai"),
            ("6", "Dashboard", "Selesai"),
            ("7", "Offline & Online", "Selesai*"),
            ("8", "Manajemen User", "Selesai"),
            ("9", "Kelas Online", "Selesai"),
            ("10", "Bank Soal", "Selesai*"),
            ("11", "Laporan", "Selesai"),
            ("12", "Cloud & Email", "Selesai*"),
        ],
    )
    note = slide.shapes.add_textbox(Inches(0.65), Inches(6.55), Inches(12), Inches(0.35))
    note.text_frame.text = "* Selesai dengan catatan teknis minor — dijelaskan di slide catatan"
    note.text_frame.paragraphs[0].font.size = Pt(10)
    note.text_frame.paragraphs[0].font.color.rgb = MUTED
    add_footer(slide, prs, 12, total_slides)

    # Module details
    module_details = [
        (
            "Modul 1–2 · Soal & Ujian",
            [
                ("📝", "Bank soal PG & esai dengan kategori mapel"),
                ("⏱", "Timer, jadwal, anti-cheat, acak urutan soal"),
                ("📊", "Quiz kelas + ujian formal lintas mapel"),
                ("🔍", "Review submissions oleh staff"),
            ],
        ),
        (
            "Modul 3–4 · Pendaftaran & Tryout",
            [
                ("📋", "Multi-tahap: Admin → Psiko → Kesehatan → Fisik"),
                ("🔒", "Berkas privat — hanya akses autentikasi"),
                ("💳", "Alur Kelas Online/Ujian dengan konfirmasi bayar"),
                ("🎯", "Tryout publik gratis tanpa login"),
            ],
        ),
        (
            "Modul 5–6 · Penilaian & Dashboard",
            [
                ("⚡", "Auto-scoring PG skala 0–100"),
                ("📈", "Grafik perkembangan real-time"),
                ("📅", "Laporan harian otomatis setelah quiz/ujian"),
                ("👥", "Dashboard khusus per role"),
            ],
        ),
        (
            "Modul 7–9 · User & Kelas",
            [
                ("📱", "PWA installable di mobile"),
                ("🏫", "Ruang kelas digital + materi per sesi"),
                ("📥", "Import user massal via CSV"),
                ("⏳", "Masa aktif per program (3–12 bulan)"),
            ],
        ),
        (
            "Modul 10–12 · Laporan & Infra",
            [
                ("📄", "PDF laporan perkembangan untuk orang tua"),
                ("📆", "Ringkasan mingguan auto-generate"),
                ("🏆", "Peringkat akademik & jasmani"),
                ("☁", "Deploy Docker + backup terjadwal"),
            ],
        ),
    ]

    page = 13
    for title, items in module_details:
        slide = prs.slides.add_slide(blank)
        add_content_header(slide, title)
        add_icon_bullets(slide, items, top=1.65, col_width=5.9)
        add_footer(slide, prs, page, total_slides)
        page += 1

    add_section_divider(prs, blank, "04", "Fitur & Timeline", "Bonus fitur dan jadwal pelaksanaan")

    # Extra features
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Fitur Tambahan", "Di luar ruang lingkup minimal — value added")
    add_feature_pills(
        slide,
        [
            ("Portal Ortu", "Undangan token · pantau perkembangan anak"),
            ("Sertifikat", "Template per program · unduh PDF"),
            ("Blog Publik", "Beranda, artikel, SEO"),
            ("Notifikasi", "Pemberitahuan real-time in-app"),
            ("PWA", "Installable di perangkat mobile"),
            ("Multi-bahasa", "Toggle Indonesia / English"),
            ("Anti-cheat", "Fullscreen, blok copy, deteksi tab"),
        ],
    )
    add_footer(slide, prs, 19, total_slides)

    # Timeline
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Timeline Pelaksanaan", "Milestone proyek Mei – Juli 2026")
    add_timeline_visual(slide)
    add_footer(slide, prs, 20, total_slides)

    # Demo flows
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Demo Alur Pengguna")
    for i, (title, steps, color) in enumerate(
        [
            (
                "Peserta",
                ["Daftar & verifikasi email", "Pendaftaran multi-tahap", "Akses kelas & quiz", "Ujian lintas mapel", "Lihat dashboard"],
                ROLE_COLORS["peserta"],
            ),
            (
                "Mentor & Admin",
                ["Kelola bank soal & tes", "Atur kelas & materi", "Review pendaftaran", "Input jasmani & laporan", "Terbitkan sertifikat"],
                ROLE_COLORS["admin"],
            ),
        ]
    ):
        x = 0.65 + i * 6.1
        card = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.65), Inches(5.85), Inches(4.8), WHITE, color, radius=0.08
        )
        hdr = add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.65), Inches(5.85), Inches(0.65), color)
        hdr.line.fill.background()
        ht = slide.shapes.add_textbox(Inches(x + 0.3), Inches(1.78), Inches(5.2), Inches(0.4))
        ht.text_frame.paragraphs[0].text = title
        ht.text_frame.paragraphs[0].font.size = Pt(20)
        ht.text_frame.paragraphs[0].font.bold = True
        ht.text_frame.paragraphs[0].font.color.rgb = WHITE

        for j, step in enumerate(steps):
            sy = 2.55 + j * 0.75
            step_card = add_shape(
                slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.25), Inches(sy), Inches(5.35), Inches(0.58), SOFT_BLUE, color, radius=0.15
            )
            num = add_shape(
                slide, MSO_SHAPE.OVAL, Inches(x + 0.4), Inches(sy + 0.1), Inches(0.38), Inches(0.38), color
            )
            set_text(num, str(j + 1), size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            st = slide.shapes.add_textbox(Inches(x + 0.9), Inches(sy + 0.12), Inches(4.5), Inches(0.35))
            st.text_frame.paragraphs[0].text = step
            st.text_frame.paragraphs[0].font.size = Pt(13)
            st.text_frame.paragraphs[0].font.color.rgb = TEXT

    add_footer(slide, prs, 21, total_slides)

    # Portal ortu + catatan combined visually
    slide = prs.slides.add_slide(blank)
    add_content_header(slide, "Portal Orang Tua & Catatan")
    ortu = add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.55), Inches(7.5), Inches(4.85), WHITE, ROLE_COLORS["parent"], radius=0.08
    )
    ortu_hdr = slide.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(6.8), Inches(0.4))
    ortu_hdr.text_frame.paragraphs[0].text = "Portal Orang Tua"
    ortu_hdr.text_frame.paragraphs[0].font.size = Pt(20)
    ortu_hdr.text_frame.paragraphs[0].font.bold = True
    ortu_hdr.text_frame.paragraphs[0].font.color.rgb = ROLE_COLORS["parent"]
    for j, line in enumerate(
        [
            "Undangan via token dari admin/mentor",
            "Login & lihat anak terhubung",
            "Grafik nilai akademik (1–100) & jasmani",
            "Unduh PDF laporan perkembangan",
        ]
    ):
        y = 2.45 + j * 0.72
        row = add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(y), Inches(6.9), Inches(0.55), SOFT_GOLD, ACCENT, radius=0.2
        )
        rt = slide.shapes.add_textbox(Inches(1.15), Inches(y + 0.1), Inches(6.5), Inches(0.35))
        rt.text_frame.paragraphs[0].text = line
        rt.text_frame.paragraphs[0].font.size = Pt(13)
        rt.text_frame.paragraphs[0].font.color.rgb = TEXT

    note_card = add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.45), Inches(1.55), Inches(4.2), Inches(4.85), RGBColor(0xFE, 0xF2, 0xF2), RGBColor(0xFC, 0xA5, 0xA5), radius=0.08
    )
    nh = slide.shapes.add_textbox(Inches(8.7), Inches(1.8), Inches(3.7), Inches(0.4))
    nh.text_frame.paragraphs[0].text = "Catatan Teknis"
    nh.text_frame.paragraphs[0].font.size = Pt(16)
    nh.text_frame.paragraphs[0].font.bold = True
    nh.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
    notes = [
        "Esai dinilai manual staff",
        "Offline = tes di lokasi fisik",
        "Bank soal = pool terpusat",
        "Revisi max 5× (Pasal 5)",
    ]
    for j, n in enumerate(notes):
        nt = slide.shapes.add_textbox(Inches(8.75), Inches(2.45 + j * 0.72), Inches(3.6), Inches(0.5))
        nt.text_frame.paragraphs[0].text = f"• {n}"
        nt.text_frame.paragraphs[0].font.size = Pt(12)
        nt.text_frame.paragraphs[0].font.color.rgb = TEXT

    add_footer(slide, prs, 22, total_slides)

    # Closing
    add_closing(prs, blank)

    return prs


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    prs.save(str(OUTPUT))
    print(f"PPT berhasil dibuat: {OUTPUT}")
    print(f"Total slide: {len(prs.slides)}")


if __name__ == "__main__":
    main()
